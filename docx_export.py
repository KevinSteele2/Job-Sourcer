"""Word (.docx) versions of the resume and cover letter, made with python-docx.

Same content, order and look as templates/documents/*.html (the PDFs):
Cambria, centered name and contact lines, bold uppercase section titles with
a rule under them, real Word bullets. Kept ATS-friendly on purpose: no tables,
text boxes, headers/footers or images, which some resume parsers skip or
scramble; everything is ordinary paragraphs in reading order.

drafter.render() calls these right after the PDFs, with the font size that
fit the PDF on one page.
"""
from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

FONT = "Cambria"


def _new_document(font_size: float, margins: tuple[float, float], line_spacing: float) -> Document:
    """A Letter-size document whose Normal style is Cambria at font_size.

    margins: (top and bottom, left and right), in inches.
    """
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(margins[0])
    section.left_margin = section.right_margin = Inches(margins[1])
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(font_size)
    # Word reads the font per script; set the East Asian one too, or some
    # characters fall back to Word's default font.
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = line_spacing
    return doc


def _paragraph(doc, text: str = "", *, size: float | None = None, bold: bool = False,
               center: bool = False, before: float = 0, after: float = 0, style: str | None = None):
    p = doc.add_paragraph(style=style)
    if text:
        run = p.add_run(text)
        run.bold = bold
        if size:
            run.font.size = Pt(size)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    return p


def _bottom_rule(paragraph) -> None:
    """A thin line under the paragraph (the HTML's border-bottom).

    python-docx has no setting for paragraph borders, so this adds Word's own
    XML for one: <w:pBdr><w:bottom .../></w:pBdr> inside the paragraph's
    properties. sz is in eighths of a point (6 = 0.75pt).
    """
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "000000")):
        bottom.set(qn(key), value)
    borders.append(bottom)
    paragraph._p.get_or_add_pPr().append(borders)


def _header(doc, contact: dict) -> None:
    """Name, then the two contact lines, centered (same as the HTML header)."""
    _paragraph(doc, contact["name"], size=20, bold=True, center=True, after=2)
    lines = [[contact.get("location"), contact.get("phone"), contact.get("email")],
             [contact.get("portfolio"), contact.get("linkedin"), contact.get("github")]]
    for parts in lines:
        text = " • ".join(p for p in parts if p)
        if text:
            _paragraph(doc, text, size=10.5, center=True)


def _section_title(doc, title: str, font_size: float) -> None:
    p = _paragraph(doc, title.upper(), size=font_size, bold=True, before=9, after=4)
    _bottom_rule(p)


def _bullet(doc, text: str) -> None:
    # "List Bullet" is Word's built-in bulleted list, so parsers see real bullets.
    _paragraph(doc, text, style="List Bullet", after=2)


def _link(paragraph, url: str, text: str) -> None:
    """A clickable link at the end of the paragraph.

    python-docx has no hyperlink API, so this builds Word's XML directly: the
    URL is registered as an external relationship of the document part, and a
    <w:hyperlink r:id=...> wraps a normal run holding the visible text.
    """
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _entry(doc, bold: str, rest: str = "", link: str = "") -> None:
    """An entry line: bold name, then " | tools", then " | link" (clickable)."""
    p = _paragraph(doc, before=4, after=1)
    p.add_run(bold).bold = True
    if rest:
        p.add_run(" | ")
        p.add_run(rest).italic = True
    if link:
        p.add_run(" | ")
        _link(p, f"https://{link}", link)


def write_resume(draft: dict, contact: dict, font_size: float, path: Path) -> None:
    resume = draft["resume"]
    doc = _new_document(font_size, margins=(0.55, 0.6), line_spacing=1.1)
    _header(doc, contact)

    if resume["summary"]:
        _section_title(doc, "Summary", font_size)
        _paragraph(doc, resume["summary"])

    # Education right after the summary, from the profile's Canonical headings
    # (as in resume.html), with the GPA.
    _section_title(doc, "Education", font_size)
    education = contact.get("education_line") or resume["education"]
    if contact.get("gpa") and "GPA" not in education:
        education += f" | GPA: {contact['gpa']}"
    _bullet(doc, education)
    if resume["coursework"]:
        _bullet(doc, "Relevant Coursework: " + ", ".join(resume["coursework"]))

    _section_title(doc, "Projects", font_size)
    links = contact.get("project_links") or {}
    for p in resume["projects"]:
        _entry(doc, p["name"], p.get("tech", ""), links.get(p["name"], ""))
        for b in p["bullets"]:
            _bullet(doc, b)

    _section_title(doc, "Experience", font_size)
    for e in resume["experience"]:
        _entry(doc, e["heading"])
        for b in e["bullets"]:
            _bullet(doc, b)
    if contact.get("additional_line"):
        p = _paragraph(doc, before=4)
        p.add_run("Additional experience: ").bold = True
        p.add_run(contact["additional_line"])

    # Labeled single-column lines: Word text columns are a common ATS problem.
    _section_title(doc, "Skills", font_size)
    for s in resume["skills"]:
        p = _paragraph(doc, after=1)
        p.add_run(f"{s['label']}: ").bold = True
        p.add_run(", ".join(s["items"]))
    doc.save(path)


def write_cover_letter(draft: dict, contact: dict, font_size: float, today: str, path: Path) -> None:
    letter = draft["cover_letter"]
    doc = _new_document(font_size + 0.5, margins=(0.9, 1.0), line_spacing=1.2)
    if letter.get("show_header", True):
        _header(doc, contact)
        _bottom_rule(doc.paragraphs[-1])
        doc.paragraphs[-1].paragraph_format.space_after = Pt(18)
        _paragraph(doc, today, after=14)
    _paragraph(doc, letter["greeting"], after=11)
    for para in letter["paragraphs"]:
        _paragraph(doc, para, after=11)
    closing = _paragraph(doc, letter["closing"], before=5)
    closing.add_run().add_break()
    closing.add_run(contact["name"])
    doc.save(path)
